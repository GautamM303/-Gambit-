"""Build the judges' slide deck: bracket_gambit/media/Bracket_Gambit.pptx

    uv run python bracket_gambit/media/make_slides.py [--frames DIR] [--out FILE]
"""
import argparse
import pathlib

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parents[1]
BG, FG, DIM, ACCENT, ACCENT2, RED = (RGBColor(18, 18, 22), RGBColor(240, 240, 240), RGBColor(160, 160, 170),
                                     RGBColor(255, 176, 0), RGBColor(90, 200, 255), RGBColor(235, 80, 70))
W, H = Inches(13.333), Inches(7.5)


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    def slide(self, title=None, subtitle=None):
        s = self.prs.slides.add_slide(self.blank)
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = BG
        self.n += 1
        if title:
            self.text(s, title, 0.6, 0.35, 12.1, 0.9, size=34, bold=True, color=FG)
            bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.2), Inches(1.6), Emu(38000))
            bar.fill.solid()
            bar.fill.fore_color.rgb = ACCENT
            bar.line.fill.background()
        if subtitle:
            self.text(s, subtitle, 0.6, 1.3, 12.1, 0.6, size=16, color=DIM)
        self.text(s, f"Bracket Gambit  ·  {self.n}", 10.8, 7.0, 2.3, 0.4, size=10, color=DIM, align=PP_ALIGN.RIGHT)
        return s

    def text(self, s, txt, x, y, w, h, size=18, bold=False, color=FG, align=PP_ALIGN.LEFT, font="Segoe UI"):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        lines = txt if isinstance(txt, list) else [txt]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = line
            r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, font
        return tb

    def bullets(self, s, items, x, y, w, h, size=17, gap=6):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(gap)
            head, _, rest = item.partition("::")
            r = p.add_run()
            r.text = "▸ " + head
            r.font.size, r.font.color.rgb, r.font.name, r.font.bold = Pt(size), ACCENT if rest else FG, "Segoe UI", bool(rest)
            if rest:
                r2 = p.add_run()
                r2.text = " " + rest.strip()
                r2.font.size, r2.font.color.rgb, r2.font.name = Pt(size), FG, "Segoe UI"
        return tb

    def image(self, s, path, x, y, w=None, h=None, caption=None):
        path = pathlib.Path(path)
        if not path.exists():
            self.text(s, f"[missing image {path.name}]", x, y, 4, 0.5, size=12, color=RED)
            return
        kw = {"width": Inches(w)} if w else {"height": Inches(h)}
        pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)
        if caption:
            self.text(s, caption, x, y + pic.height / 914400 + 0.05, pic.width / 914400, 0.4, size=11, color=DIM)
        return pic

    def box(self, s, txt, x, y, w, h, fill=ACCENT, color=BG, size=14, bold=True):
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        sh.line.fill.background()
        tf = sh.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = txt
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, "Segoe UI"
        return sh

    def table(self, s, rows, x, y, w, col_w, size=13, header=True):
        n_rows, n_cols = len(rows), len(rows[0])
        shape = s.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(0.4 * n_rows))
        t = shape.table
        for j, cw in enumerate(col_w):
            t.columns[j].width = Inches(cw)
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = t.cell(i, j)
                c.fill.solid()
                c.fill.fore_color.rgb = RGBColor(40, 40, 48) if (i == 0 and header) else (RGBColor(26, 26, 32) if i % 2 else RGBColor(22, 22, 28))
                tf = c.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                r = p.add_run()
                r.text = str(val)
                r.font.size, r.font.name = Pt(size), "Segoe UI"
                r.font.bold = (i == 0 and header)
                colour = FG
                if val.startswith("✔"):
                    colour = RGBColor(120, 230, 120)
                elif val.startswith("◐"):
                    colour = ACCENT
                elif val.startswith("✖") or val.startswith("—"):
                    colour = DIM
                r.font.color.rgb = ACCENT if (i == 0 and header) else colour
        return shape


def pick_frame(frames, pattern, index=-1):
    found = sorted(frames.glob(pattern))
    return found[index] if found else frames / pattern


def build(frames: pathlib.Path, out: pathlib.Path):
    d = Deck()
    sim = ROOT / "chopped_urdf_v2" / "sim"
    game_over_png = pick_frame(frames, "*_game_over.png")
    observing_png = pick_frame(frames, "*_observing.png", -1)

    # 1 title
    s = d.slide()
    d.text(s, "Bracket Gambit", 0.8, 2.0, 11.7, 1.4, size=60, bold=True)
    d.text(s, "BracketBot plays chess against you — camera, Stockfish, and its own arm", 0.8, 3.3, 11.7, 0.8, size=26, color=ACCENT)
    d.text(s, ["A physically embodied chess opponent built on the Bracket Bot SDK (bbos).",
               "Perception → legal-move inference → Stockfish →",
               "pick-and-place → verification, every turn."],
           0.8, 4.4, 7.6, 1.6, size=18, color=DIM)
    d.image(s, game_over_png, 8.9, 4.5, w=4.0)

    # 2 the experience
    s = d.slide("What the judges see", "One turn of the game, as the robot experiences it")
    steps = [("1  Observe", "Head camera reads the board.\nSetup verified: 16 white / 16 black."),
             ("2  Human moves", "\"My turn is complete\" — Enter,\nweb button, or the wake word."),
             ("3  Infer the move", "Occupancy change matched\nagainst every legal move."),
             ("4  Validate", "Illegal / ambiguous? The robot\nsays which squares are wrong."),
             ("5  Think", "Stockfish 19 over UCI,\nadjustable skill, 1 s per move."),
             ("6  Announce", "\"I play Nf3: knight from g1 to f3.\"\n\"Why?\" → grounded explanation."),
             ("7  Move the piece", "Pick by the shaft, carry, place;\ncaptures go to a tray first."),
             ("8  Verify", "Re-read the board; the internal\nposition advances only on a match."),
             ("9  Repeat", "Check, mate and the end of the demo\nannounced; arm parks safely.")]
    for i, (h, b) in enumerate(steps):
        col, row = i % 3, i // 3
        x, y = 0.6 + col * 4.15, 1.9 + row * 1.75
        d.box(s, h, x, y, 3.9, 0.5, fill=ACCENT if row != 1 else ACCENT2)
        d.text(s, b.split("\n"), x, y + 0.55, 3.9, 1.1, size=14, color=FG)

    # 3 capabilities map
    s = d.slide("Bracket Bot capabilities → what Bracket Gambit uses", "From the handout / SDK demos (demos/view_*.py, daemon.py, constants.py)")
    rows = [["Bracket Bot capability", "SDK channel / API", "In Bracket Gambit"],
            ["Head camera", "camera.head.jpeg", "✔ board reading: ArUco corners → homography → 64 squares"],
            ["Right / left 7-DOF arm + gripper", "arm_*.state/.ctrl/.torque, cfg.ik (hybrid_ik)", "✔ right arm: IK straight-line pick & place, grasp check from the gripper angle; left arm idle"],
            ["Speaker", "speaker.audio (PCM)", "✔ announces moves, check/mate, help requests (espeak-ng); chimes fallback"],
            ["Microphone + wake word", "mic.audio, wakeword.state", "✔ wake word = \"my turn is complete\" (no speech-to-text in the SDK → other commands via keys/web)"],
            ["LED", "led.ctrl", "✔ phase colours: waiting / looking / thinking / moving / needs help"],
            ["Depth camera (point cloud)", "camera.points", "◐ used in the earlier pick-by-voice demo; chess uses the 2-D homography instead"],
            ["Navigation / base drive", "— (not in the demos we had)", "✖ robot parked by hand; the sim's roll-up is simulation-only"],
            ["IMU, Quest teleop, USB tree", "imu.*, quest.*, usb.tree", "✖ not needed"],
            ["MuJoCo simulation", "chopped_urdf_v2 (CAD-accurate URDF)", "✔ every manipulation and vision path validated there first"]]
    d.table(s, rows, 0.6, 1.9, 12.1, [3.0, 3.6, 5.5], size=12)

    # 4 perception
    s = d.slide("Perception: from a camera frame to 64 squares", "No piece recognition needed — occupancy + colour is enough to play")
    d.bullets(s, ["Corner markers:: four ArUco tags (ids 0–3) fix the board's homography every turn — any camera angle, no calibration of intrinsics",
                  "Rectify:: the board is warped to a 400×400 top-down view; each square is sampled in a band at its centre (where the piece base sits)",
                  "Reference diff:: an empty-board photo taken at calibration; a square is occupied if it differs — works for any board or piece colours",
                  "Colour learned, not hard-coded:: the black/white split is learned from the 32 pieces of the starting position (Otsu)",
                  "Robustness found by testing:: a tall far-rank piece smears into the square behind it; the sample band was chosen by sweeping rendered positions (0 errors vs 4–11)"],
              0.6, 1.9, 6.3, 5.0, size=15)
    d.image(s, observing_png, 7.1, 1.9, w=5.7, caption="Right: head camera with the detected grid; bottom: internal position, dots = what the camera saw")

    # 5 understanding the move
    s = d.slide("Understanding the human's move", "Legal-move matching instead of tracking hands")
    d.bullets(s, ["Diff against legal moves:: for every legal move, predict the occupancy and count mismatching squares; accept only a unique zero-mismatch match",
                  "Captures, castling, en passant:: fall out naturally — the occupancy pattern identifies them",
                  "Rejection with reasons:: \"c6 should be empty but I see a black piece\" — the robot waits for a corrected board instead of guessing",
                  "No change / hand in view:: detected and reported; the camera view is re-read on the next signal",
                  "\"What did you see?\":: the robot repeats the counts and its last inference on request"],
              0.6, 1.9, 6.3, 5.0, size=15)
    d.image(s, HERE / "rejected_move.png", 7.1, 1.9, w=5.7, caption="A pawn moved to an impossible square: the closest legal move and the offending squares (red) are reported")

    # 6 brain
    s = d.slide("The brain: Stockfish, kept honest", "Strength you can dial; explanations grounded in the position")
    d.bullets(s, ["Stockfish 19 over UCI:: python-chess, Skill Level 0–20, time per move configurable; the same engine the strongest programs use",
                  "Physical constraints fed back:: promotions are not executable by the gripper → the best non-promoting move is chosen instead",
                  "\"Why did you play that?\":: one or two sentences built from facts of the move (capture, check, development, centre, threats) plus Stockfish's evaluation and principal variation — never invented intent",
                  "Labelled fallback:: if the binary is missing a small built-in searcher plays, and every status line says \"fallback (NOT Stockfish)\""],
              0.6, 1.9, 7.6, 4.5, size=15)
    d.box(s, "\"I played this because it takes a pawn. Stockfish rates the position +1.6 pawns for me, expecting 3. cxd6 e5 4. Nc3.\"",
          8.5, 2.2, 4.3, 2.2, fill=RGBColor(34, 34, 44), color=FG, size=14, bold=False)
    d.box(s, "\"I play Nf3: knight from g1 to f3.\"", 8.5, 4.7, 4.3, 0.9, fill=RGBColor(34, 34, 44), color=ACCENT2, size=14, bold=False)

    # 7 manipulation
    s = d.slide("Manipulation: deliberate, verified, recoverable", "Validated in MuJoCo with the robot's CAD-accurate URDF")
    d.bullets(s, ["Pick by the shaft:: approach from above, pads oriented diagonally so the claws clear the neighbours; lift and carry at conservative speeds",
                  "Captures first:: the captured piece goes to a tray beside the board before the attacker moves; castling moves king then rook",
                  "Grasp verification:: sim — lift height + contact; real arm — gripper angle after closing (\"did I actually get it?\")",
                  "Retry, then ask:: a miss is retried at the position the camera saw the piece; a second miss parks the arm and asks the human to finish the step, then re-verifies",
                  "Never trust, always look:: the chess position is updated only after the camera confirms the physical board"],
              0.6, 1.9, 6.3, 5.0, size=15)
    d.image(s, game_over_png, 7.1, 1.9, w=5.7, caption="End of the recorded game: two black pieces carried to the tray (left of the board); every placement within 1 mm")

    # 8 safety & control
    s = d.slide("Safety and control", "Fail safely around people, the board and the hardware")
    phases = ["IDLE", "OBSERVING", "WAITING", "THINKING", "ANNOUNCING", "MOVING", "VERIFYING", "RECOVERING", "NEEDS HELP", "PAUSED", "STOPPED"]
    for i, p in enumerate(phases):
        colour = RED if p in ("NEEDS HELP", "STOPPED") else (ACCENT2 if p in ("RECOVERING", "PAUSED") else ACCENT)
        d.box(s, p, 0.6 + (i % 6) * 2.05, 1.95 + (i // 6) * 0.7, 1.9, 0.5, fill=colour, size=12)
    d.bullets(s, ["Emergency stop:: 'e' in the terminal, the red web button, or Ctrl-C sets a flag polled every 20 ms control tick; torque is cut on all joints; the daemon also cuts torque if the controller disappears",
                  "Before any motion:: the whole straight line is solved by IK and checked against a workspace box — an unreachable target never leaves the arm half-way",
                  "Three control surfaces:: terminal keys · phone-friendly web page (port 8010, doubles as spectator display) · wake word — voice is never a dependency",
                  "Status everywhere:: LED colour per phase, spoken status, live display with board map, camera overlay, engine line and event log",
                  "Dry run first:: on the robot, `play --robot` prints every waypoint and moves nothing until `--execute`"],
              0.6, 3.5, 12.1, 3.6, size=14)

    # 9 results
    s = d.slide("Results and validation", "What has actually been run")
    d.table(s, [["Evidence", "Result"],
                ["Simulated game (this deck's video), Stockfish 19", "1. e4 d5 2. f3 e5 3. Qe2 Nc6 4. exd5 Nf6 5. dxc6 — two captures to the tray, every human move inferred with 0 mismatches, placements ≤ 1 mm, 0 illegal contacts"],
                ["Earlier 6-move simulated game", "two captures to the tray, 0 collisions, 0 helper resets"],
                ["Vision fixtures (real sim renders, 7 positions)", "64/64 squares correct in every position, every move inferred"],
                ["Unit + integration tests", "30 tests: vision, engine, manipulation planning/retries, full loop with a mock robot (illegal move, no change, failed grasp → help, pause/e-stop)"],
                ["Web UI / e-stop path", "exercised end-to-end against the mock game"],
                ["Real robot", "backend written against the daemon sources; NOT yet run on hardware (see next slide)"]],
            0.6, 1.9, 12.1, [4.6, 7.5], size=13)

    # 10 honesty
    s = d.slide("What is real, simulated, and still to verify", "We would rather tell you than have you find out")
    d.table(s, [["", "Status"],
                ["Vision, move inference, Stockfish, state machine, retries, UI", "✔ implemented and tested (synthetic + sim renders + mock robot)"],
                ["Arm pick-and-place, captures, verification", "✔ verified in MuJoCo with the robot's URDF; real-arm backend uses the same sequence through the daemon's IK"],
                ["Base rolling up to the table", "◐ simulation only — no base-drive channel in the SDK we had; robot is parked by hand"],
                ["Speech, wake word, LED, e-stop on hardware", "◐ implemented on the documented channels; hardware-unverified"],
                ["Promotion, arbitrary start positions, gestures with the arms, left arm", "✖ omitted on purpose to protect the core loop"],
                ["Feasibility on the physical Bracket Bot", "reach ≈ 24 cm deep → 3 cm squares; ~3.9 Nm shoulder vs 30 g pieces fine; main risk: the self-balancing base leaning as the arm extends → prop it for the demo"]],
            0.6, 1.9, 12.1, [4.6, 7.5], size=13)

    # 11 roadmap
    s = d.slide("Next steps", "In priority order")
    d.bullets(s, ["1  Bring-up on the robot:: check → calibrate (reference image, teach a1/h1/a8, gripper angles) → dry run → execute; 7-item hardware checklist in STATUS.md",
                  "2  Grip current as a second grasp signal:: detect a slipping piece during the carry",
                  "3  Camera-refined pick position on every move:: today only on the retry",
                  "4  Left arm for the a/b files:: if the parked pose cannot reach them; the interface already routes each transfer through pick/place",
                  "5  Expressive gestures:: look at the player, a thinking sway — behind a flag, verified separately from manipulation",
                  "6  Castling and en passant on the physical board, promotion via a piece swap, resume after an e-stop, coaching mode"],
              0.6, 1.9, 12.1, 5.0, size=16)

    # 12 closing
    s = d.slide()
    d.text(s, "Bracket Gambit", 0.8, 1.6, 11.7, 1.0, size=48, bold=True)
    d.text(s, "Play a game against the robot.", 0.8, 2.6, 11.7, 0.8, size=28, color=ACCENT)
    d.text(s, ["uv run python gambit.py play --sim --moves 4        # laptop, MuJoCo",
               "uv run gambit.py play --robot --execute --moves 6   # on the robot, after calibrate",
               "",
               "Code: gambit.py + bracket_gambit/   ·   Docs: README, DEMO_SCRIPT, STATUS   ·   Tests: uv run pytest tests"],
           0.8, 3.7, 11.7, 2.0, size=16, color=DIM, font="Consolas")
    d.text(s, "Built on the Bracket Bot SDK (bbos), MuJoCo, OpenCV, python-chess and Stockfish 19.", 0.8, 6.2, 11.7, 0.6, size=14, color=DIM)

    out.parent.mkdir(parents=True, exist_ok=True)
    d.prs.save(str(out))
    print(f"wrote {out} ({d.n} slides)")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--frames", type=pathlib.Path, default=ROOT / "chopped_urdf_v2" / "sim" / "gambit_frames")
    ap.add_argument("--out", type=pathlib.Path, default=HERE / "Bracket_Gambit.pptx")
    args = ap.parse_args()
    build(args.frames, args.out)
